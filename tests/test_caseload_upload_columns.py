"""Caseload upload: columns matched by name, the "English Learner" column
(blank = not an EL), EL / gender value normalization, and Gender in the
template and export.

A counselor's caseload can be every English Learner plus a slice of general
ed (e.g. last names V-Z), exported from the SIS with a column titled
"English Learner" that is blank for the general-ed students. That file must
import as-is: the blank cells are EO, not errors, and the column is found by
its name wherever it sits.
"""
import io

import pytest
from flask import url_for
from openpyxl import Workbook, load_workbook

from app import db
from app.models.student import Student
from app.models.user import User


def _xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _headers_of(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    row = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return [str(v or '').strip() for v in row]


def _scrub():
    Student.query.filter(Student.student_id_number.like('COL-%')).delete(
        synchronize_session=False)
    User.query.filter_by(username='col_me').delete(synchronize_session=False)
    db.session.commit()


@pytest.fixture
def col_env(app):
    """One counselor with an existing Newcomer student, logged in."""
    with app.app_context():
        db.session.rollback()
        _scrub()
        me = User(username='col_me', display_name='Col Me', role='counselor',
                  setup_completed=True)
        me.set_password('passw0rd123')
        db.session.add(me)
        db.session.commit()
        existing = Student(student_id_number='COL-EXIST', first_name='Elena',
                           last_name='Vargas', grade_level=9, status='active',
                           assigned_counselor_id=me.id, el_status='Newcomer',
                           el_level='EL 2', ell_status=True, gender='Female',
                           email='elena@example.org')
        db.session.add(existing)
        db.session.commit()
        ids = dict(me=me.id, existing=existing.id)
    client = app.test_client()
    with client.session_transaction() as sess:
        sess['_user_id'] = str(ids['me'])
        sess['_fresh'] = True
    yield client, ids
    with app.app_context():
        db.session.rollback()
        _scrub()


def _upload(client, headers, rows):
    return client.post('/caseload/upload',
                       data={'file': (_xlsx(headers, rows), 'roster.xlsx')},
                       content_type='multipart/form-data', follow_redirects=False)


def _preview(client, headers, rows):
    return client.post('/caseload/upload/preview',
                       data={'file': (_xlsx(headers, rows), 'roster.xlsx')},
                       content_type='multipart/form-data')


def _student(app, sid):
    with app.app_context():
        return Student.query.filter_by(student_id_number=sid).first()


# ── the "English Learner" column ───────────────────────────────────────────

def test_english_learner_column_blank_means_not_an_el(app, col_env):
    client, ids = col_env
    # SIS layout: different order, ID first, column titled "English Learner".
    headers = ['Student ID #', 'Last Name', 'First Name', 'Grade', 'English Learner']
    r = _upload(client, headers, [
        ('COL-1', 'Vega', 'Ana', 10, ''),        # general-ed (V-Z): blank
        ('COL-2', 'Ybarra', 'Luis', 11, 'LTEL'),
        ('COL-3', 'Zamora', 'Eva', 9, 'Yes'),    # "is an EL", no level given
    ])
    assert r.status_code in (200, 302)
    ana, luis, eva = (_student(app, s) for s in ('COL-1', 'COL-2', 'COL-3'))
    assert ana.el_status == 'EO' and ana.ell_status is False
    assert luis.el_status == 'LTEL' and luis.ell_status is True
    assert eva.el_status == 'LTEL' and eva.ell_status is True


def test_generic_yes_with_a_level_is_a_newcomer(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #',
               'English Learner', 'EL Level']
    _upload(client, headers, [
        ('Ana', 'Vega', 9, 'COL-4', 'Y', 'EL 1'),
        ('Luis', 'Ybarra', 9, 'COL-5', 'yes', ''),
        ('Eva', 'Zamora', 9, 'COL-6', 'X', 'el2'),
    ])
    got = {s: (_student(app, s).el_status, _student(app, s).el_level)
           for s in ('COL-4', 'COL-5', 'COL-6')}
    assert got == {'COL-4': ('Newcomer', 'EL 1'), 'COL-5': ('LTEL', ''),
                   'COL-6': ('Newcomer', 'EL 2')}


def test_el_status_spellings_are_normalized(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'EL Status']
    _upload(client, headers, [
        ('A', 'A', 9, 'COL-7', 'newcomer'),
        ('B', 'B', 9, 'COL-8', 'rfep'),
        ('C', 'C', 9, 'COL-9', 'IFEP'),           # never an EL → EO
        ('D', 'D', 9, 'COL-10', 'English Only'),
        ('E', 'E', 9, 'COL-11', 'No'),
    ])
    got = {s: _student(app, s).el_status
           for s in ('COL-7', 'COL-8', 'COL-9', 'COL-10', 'COL-11')}
    assert got == {'COL-7': 'Newcomer', 'COL-8': 'RFEP', 'COL-9': 'EO',
                   'COL-10': 'EO', 'COL-11': 'EO'}


def test_unrecognized_el_value_is_a_row_error_not_a_guess(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'English Learner']
    data = _preview(client, headers, [('A', 'A', 9, 'COL-12', 'Maybe')]).get_json()
    assert data['ok'] is True
    assert any('Invalid EL Status' in e for e in data['errors'])
    _upload(client, headers, [('A', 'A', 9, 'COL-12', 'Maybe')])
    assert _student(app, 'COL-12') is None


def test_preview_explains_inferred_el_statuses(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'English Learner']
    data = _preview(client, headers, [('Eva', 'Zamora', 9, 'COL-13', 'Yes')]).get_json()
    assert data['counts']['notices'] == len(data['notices']) >= 1
    assert any('LTEL' in n and 'Zamora, Eva' in n for n in data['notices'])


# ── columns by name ────────────────────────────────────────────────────────

def test_missing_required_column_is_a_clear_fatal(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Student ID #', 'English Learner']  # no Grade
    r = _preview(client, headers, [('A', 'A', 'COL-14', '')])
    assert r.status_code == 400
    err = r.get_json()['error']
    assert 'Required column' in err and 'Grade' in err


def test_absent_optional_columns_leave_existing_values_alone(app, col_env):
    client, ids = col_env
    # Only the required columns: the Newcomer keeps status/level/gender/email.
    _upload(client, ['First Name', 'Last Name', 'Grade', 'Student ID #'],
            [('Elena', 'Vargas', 10, 'COL-EXIST')])
    s = _student(app, 'COL-EXIST')
    assert s.grade_level == 10
    assert (s.el_status, s.el_level, s.ell_status) == ('Newcomer', 'EL 2', True)
    assert s.gender == 'Female' and s.email == 'elena@example.org'


def test_blank_el_cell_in_a_present_column_resets_to_eo(app, col_env):
    """Blank is meaningful when the column exists: the roster says this
    student is not an English Learner."""
    client, ids = col_env
    _upload(client, ['First Name', 'Last Name', 'Grade', 'Student ID #', 'English Learner'],
            [('Elena', 'Vargas', 9, 'COL-EXIST', '')])
    s = _student(app, 'COL-EXIST')
    assert (s.el_status, s.el_level, s.ell_status) == ('EO', '', False)


def test_blank_gender_and_email_cells_keep_existing_values(app, col_env):
    client, ids = col_env
    _upload(client, ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Gender', 'Email'],
            [('Elena', 'Vargas', 9, 'COL-EXIST', '', '')])
    s = _student(app, 'COL-EXIST')
    assert s.gender == 'Female' and s.email == 'elena@example.org'


def test_legacy_template_layout_still_imports(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Email',
               'EL Status', 'EL Level', 'IEP', '504 Plan']
    _upload(client, headers, [('A', 'A', 9, 'COL-15', 'a@x.org', 'EO', '', 'Yes', '')])
    s = _student(app, 'COL-15')
    assert s.el_status == 'EO' and s.iep_status is True and s.section_504 is False


# ── gender ─────────────────────────────────────────────────────────────────

def test_gender_column_is_read_and_normalized(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Gender']
    _upload(client, headers, [
        ('A', 'A', 9, 'COL-16', 'M'),
        ('B', 'B', 9, 'COL-17', 'female'),
        ('C', 'C', 9, 'COL-18', 'X'),
        ('D', 'D', 9, 'COL-19', 'Non-Binary'),
    ])
    got = {s: _student(app, s).gender for s in ('COL-16', 'COL-17', 'COL-18', 'COL-19')}
    assert got == {'COL-16': 'Male', 'COL-17': 'Female', 'COL-18': 'Non-Binary',
                   'COL-19': 'Non-Binary'}


def test_template_and_export_carry_gender(app, col_env):
    client, ids = col_env
    with app.test_request_context():
        tpl = url_for('caseload.download_template')
        exp = url_for('caseload.export_caseload')
    tpl_headers = _headers_of(client.get(tpl).data)
    assert 'Gender' in tpl_headers and 'EL Status' in tpl_headers
    exp_headers = _headers_of(client.get(exp).data)
    assert 'Gender' in exp_headers
    # Round trip: the export's own headers import cleanly.
    r = _preview(client, exp_headers, [])
    assert r.status_code == 200 and r.get_json()['ok'] is True


# ── file name handling ─────────────────────────────────────────────────────

def test_uppercase_xlsx_extension_is_accepted(app, col_env):
    """Windows/Excel save 'MASTER CASELOAD 26-27.XLSX'; the extension check
    must not be case-sensitive."""
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'English Learner']
    r = client.post('/caseload/upload/preview',
                    data={'file': (_xlsx(headers, [('Ana', 'Vega', 10, 'COL-20', '')]),
                                   'MASTER CASELOAD 26-27.XLSX')},
                    content_type='multipart/form-data')
    assert r.status_code == 200, r.get_json()
    data = r.get_json()
    assert data['ok'] is True
    assert [n['sid'] for n in data['new']] == ['COL-20']


# ── the SIS export's "Idea Cur Level" column ───────────────────────────────

def test_idea_cur_level_column_with_sis_wording(app, col_env):
    """The district export titles the language-program field "Idea Cur Level"
    and fills it with English Learners / Redesignated FEP / English Only /
    IFEP. All four must land on the app's statuses without any renaming."""
    client, ids = col_env
    headers = ['Student ID #', 'Last Name', 'First Name', 'Grade', 'Idea Cur Level']
    r = _upload(client, headers, [
        ('COL-21', 'Vega', 'Ana', 10, 'English Learners'),
        ('COL-22', 'Ybarra', 'Luis', 11, 'Redesignated FEP'),
        ('COL-23', 'Zamora', 'Eva', 9, 'English Only'),
        ('COL-24', 'Zuniga', 'Omar', 12, 'IFEP'),
        ('COL-25', 'Villa', 'Rosa', 9, ''),          # blank → not an EL
    ])
    assert r.status_code in (200, 302)
    got = {s: (_student(app, s).el_status, _student(app, s).ell_status)
           for s in ('COL-21', 'COL-22', 'COL-23', 'COL-24', 'COL-25')}
    assert got == {
        'COL-21': ('LTEL', True),      # current EL, no level given → LTEL
        'COL-22': ('RFEP', True),
        'COL-23': ('EO', False),
        'COL-24': ('EO', False),       # initially fluent — never an EL
        'COL-25': ('EO', False),
    }


def test_idea_cur_level_english_learners_with_a_level_is_newcomer(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #',
               'Idea Cur Level', 'EL Level']
    data = _preview(client, headers,
                    [('Ana', 'Vega', 9, 'COL-26', 'English Learners', 'EL 1'),
                     ('Luis', 'Ybarra', 9, 'COL-27', 'English Learners', '')]).get_json()
    assert data['ok'] is True and not data['errors']
    # Only the level-less one is an inference worth flagging.
    assert any('LTEL' in n and 'Ybarra, Luis' in n for n in data['notices'])
    _upload(client, headers,
            [('Ana', 'Vega', 9, 'COL-26', 'English Learners', 'EL 1'),
             ('Luis', 'Ybarra', 9, 'COL-27', 'English Learners', '')])
    assert (_student(app, 'COL-26').el_status, _student(app, 'COL-26').el_level) == ('Newcomer', 'EL 1')
    assert _student(app, 'COL-27').el_status == 'LTEL'


# ── hand-entered EL Levels survive a roster refresh ────────────────────────

def test_roster_refresh_keeps_hand_entered_newcomer_level(app, col_env):
    """'English Learners' only says the student IS an EL. The counselor hand-
    enters EL Levels for newcomers, so re-uploading the same roster must not
    downgrade them back to LTEL."""
    client, ids = col_env                      # COL-EXIST is Newcomer / EL 2
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Idea Cur Level']
    _upload(client, headers, [('Elena', 'Vargas', 10, 'COL-EXIST', 'English Learners')])
    s = _student(app, 'COL-EXIST')
    assert (s.el_status, s.el_level, s.ell_status) == ('Newcomer', 'EL 2', True)
    assert s.grade_level == 10                 # the rest of the row still applied


def test_explicit_status_in_the_file_still_wins(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Idea Cur Level']
    _upload(client, headers, [('Elena', 'Vargas', 9, 'COL-EXIST', 'Redesignated FEP')])
    s = _student(app, 'COL-EXIST')
    assert (s.el_status, s.el_level, s.ell_status) == ('RFEP', '', True)


def test_generic_marker_still_upgrades_a_non_el_to_ltel(app, col_env):
    client, ids = col_env
    headers = ['First Name', 'Last Name', 'Grade', 'Student ID #', 'Idea Cur Level']
    _upload(client, headers, [('Ana', 'Vega', 9, 'COL-28', '')])                 # EO first
    assert _student(app, 'COL-28').el_status == 'EO'
    _upload(client, headers, [('Ana', 'Vega', 9, 'COL-28', 'English Learners')])
    assert (_student(app, 'COL-28').el_status, _student(app, 'COL-28').ell_status) == ('LTEL', True)
